#!/usr/bin/env python3
"""
Универсальный парсер: извлекает из файлов
  - номера банковских карт (с определением BIN / эмитента),
  - телефонные номера (с определением страны),
  - учётные записи Telegram,
  - сообщения (переписки) из экспорта Telegram Desktop,
и формирует Excel с отдельными листами + сводкой.
"""

import argparse
import csv
import io
import json
import logging
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional

import pandas as pd
import phonenumbers
from phonenumbers import region_code_for_number
from phonenumbers.geocoder import country_name_for_number

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
log = logging.getLogger("parser")

SUPPORTED_EXTS = {
    ".txt", ".log", ".md", ".ini", ".conf", ".rst",
    ".csv", ".json", ".xlsx", ".docx", ".pdf", ".html", ".htm",
}


# ===========================================================================
# 1. ЧТЕНИЕ ФАЙЛОВ
# ===========================================================================
def read_text_auto(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_bytes().decode("utf-8", errors="ignore")


def read_csv_text(path: Path) -> str:
    text = read_text_auto(path)
    parts = []
    try:
        for row in csv.reader(io.StringIO(text)):
            parts.extend(c for c in row if c)
    except csv.Error:
        return text
    return "\n".join(parts)


def read_json_text(path: Path) -> str:
    data = json.loads(read_text_auto(path))
    parts = []

    def walk(o):
        if isinstance(o, dict):
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)
        elif isinstance(o, (str, int, float)):
            parts.append(str(o))

    walk(data)
    return "\n".join(parts)


def read_xlsx_text(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                parts.extend(str(v) for v in row if v is not None)
    finally:
        wb.close()
    return "\n".join(parts)


def read_docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def read_pdf_text(path: Path) -> str:
    import pdfplumber
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    return "\n".join(parts)


def read_html_text(path: Path) -> str:
    text = read_text_auto(path)
    text = re.sub(r"<script.*?</script>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<style.*?</style>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("&nbsp;", " ").replace("&amp;", "&")
    return re.sub(r"[ \t]+", " ", text)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".log", ".md", ".ini", ".conf", ".rst"}:
        return read_text_auto(path)
    if ext == ".csv":
        return read_csv_text(path)
    if ext == ".json":
        return read_json_text(path)
    if ext == ".xlsx":
        return read_xlsx_text(path)
    if ext == ".docx":
        return read_docx_text(path)
    if ext == ".pdf":
        return read_pdf_text(path)
    if ext in {".html", ".htm"}:
        return read_html_text(path)
    return read_text_auto(path)


# ===========================================================================
# 2. СПРАВОЧНИК BIN
# ===========================================================================
class BinDatabase:
    def __init__(self, csv_path: Optional[str] = None):
        self.by_6: dict[str, dict] = {}
        self.by_8: dict[str, dict] = {}
        self.loaded = False
        if csv_path and Path(csv_path).exists():
            self._load(Path(csv_path))

    def _load(self, path: Path):
        try:
            df = pd.read_csv(
                path, dtype=str, keep_default_na=False,
                usecols=lambda c: c in {
                    "BIN", "Brand", "Type", "Category",
                    "Issuer", "CountryName", "isoCode2", "isoCode3"
                },
            )
        except Exception as e:
            log.error("Не удалось загрузить BIN-справочник %s: %s", path, e)
            return

        col = {c.lower(): c for c in df.columns}
        bin_col = col.get("bin")
        if not bin_col:
            log.error("В CSV нет колонки BIN")
            return

        for _, row in df.iterrows():
            bin_val = str(row[bin_col]).strip().zfill(6)
            rec = {
                "brand":    row.get(col.get("brand", ""), "") or "",
                "type":     row.get(col.get("type", ""), "") or "",
                "category": row.get(col.get("category", ""), "") or "",
                "issuer":   row.get(col.get("issuer", ""), "") or "",
                "country":  row.get(col.get("countryname", ""), "") or "",
                "iso2":     row.get(col.get("isocode2", ""), "") or "",
                "iso3":     row.get(col.get("isocode3", ""), "") or "",
            }
            if len(bin_val) >= 8:
                self.by_8[bin_val[:8]] = rec
            self.by_6[bin_val[:6]] = rec

        self.loaded = True
        log.info("BIN-справочник загружен: %d записей (6-значных ключей)",
                 len(self.by_6))

    def lookup(self, card_digits: str) -> dict:
        if not self.loaded:
            return {}
        if len(card_digits) >= 8 and card_digits[:8] in self.by_8:
            return self.by_8[card_digits[:8]]
        return self.by_6.get(card_digits[:6], {})


# ===========================================================================
# 3. ПОИСК КАРТ
# ===========================================================================
CARD_RE = re.compile(r"(?<!\d)(?:\d[ \t\-]?){12,18}\d(?!\d)")


def luhn_ok(number: str) -> bool:
    digits = [int(c) for c in number]
    n = len(digits)
    total = 0
    for i, d in enumerate(digits):
        if (n - i) % 2 == 0:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def card_brand_by_prefix(num: str) -> str:
    if num[:2] in ("34", "37"):                       return "American Express"
    if num[0] == "4":                                 return "Visa"
    if 2221 <= int(num[:4]) <= 2720:                  return "MasterCard"
    if 51 <= int(num[:2]) <= 55:                      return "MasterCard"
    if 2200 <= int(num[:4]) <= 2204:                  return "Мир"
    if 3528 <= int(num[:4]) <= 3589:                  return "JCB"
    if num[:2] == "62":                               return "UnionPay"
    if num[:2] == "50" or 56 <= int(num[:2]) <= 69:   return "Maestro"
    if num[:4] == "6011" or num[:2] == "65":          return "Discover"
    if 300 <= int(num[:3]) <= 305 or num[:2] in ("36", "38", "39"):
        return "Diners Club"
    return "Неизвестно"


# ===========================================================================
# 4. ПОИСК ТЕЛЕФОНОВ
# ===========================================================================
PHONE_RE = re.compile(r"(?<!\d)(\+?\d[\d\s\-\(\)]{7,20}\d)(?!\d)")


def detect_country_full(digits: str) -> tuple:
    """Возвращает (ISO-код, название страны на русском) или ('', '')."""
    try:
        number = phonenumbers.parse("+" + digits, None)
        iso = region_code_for_number(number) or ""
        name = country_name_for_number(number, "ru") or ""
        return iso, name
    except phonenumbers.NumberParseException:
        return "", ""


def format_phone(digits: str) -> str:
    try:
        number = phonenumbers.parse("+" + digits, None)
        if phonenumbers.is_valid_number(number):
            return phonenumbers.format_number(
                number, phonenumbers.PhoneNumberFormat.INTERNATIONAL
            )
    except phonenumbers.NumberParseException:
        pass
    return "+" + digits if len(digits) > 10 else digits


# ===========================================================================
# 5. СТРУКТУРЫ ДАННЫХ
# ===========================================================================
@dataclass
class CardMatch:
    value: str
    brand: str
    card_type: str
    category: str
    issuer: str
    issuer_country: str
    bin_code: str
    luhn: bool
    raw: str
    source_file: str
    context: str


@dataclass
class PhoneMatch:
    value: str
    country: str
    raw: str
    source_file: str
    context: str


@dataclass
class TgAccount:
    user_id: str
    username: str
    first_name: str
    last_name: str
    phone: str
    is_bot: str
    is_premium: str
    source_chat: str
    source_file: str


@dataclass
class TgMessage:
    msg_id: str
    chat_name: str
    chat_type: str
    sender_id: str
    sender_name: str
    sender_username: str
    date: str
    time: str
    datetime_iso: str
    msg_type: str
    text: str
    reply_to: str
    edited: str
    source_file: str


@dataclass
class UniqueCard:
    value: str
    digits: str
    bin_code: str
    brand: str
    card_type: str
    issuer: str
    country: str
    occurrences: int = 0
    files: set = field(default_factory=set)


@dataclass
class UniquePhone:
    value: str
    digits: str
    country_iso: str
    country_name: str
    occurrences: int = 0
    files: set = field(default_factory=set)


def snippet(text: str, start: int, end: int, pad: int = 40) -> str:
    lo, hi = max(0, start - pad), min(len(text), end + pad)
    return text[lo:hi].replace("\n", " ").replace("\r", " ").strip()


# ===========================================================================
# 6. ОБРАБОТКА ОБЫЧНОГО ФАЙЛА (карты + телефоны)
# ===========================================================================
def process_file(path: Path, bindb: BinDatabase):
    try:
        text = extract_text(path)
    except Exception as e:
        log.error("Ошибка чтения %s: %s", path, e)
        return [], []

    cards, phones = [], []
    seen = set()

    for m in CARD_RE.finditer(text):
        raw = m.group(0)
        digits = re.sub(r"\D", "", raw)
        if not 13 <= len(digits) <= 19:
            continue
        luhn = luhn_ok(digits)
        bin_rec = bindb.lookup(digits)
        if not luhn and not bin_rec:
            continue
        key = ("c", digits)
        if key in seen:
            continue
        seen.add(key)
        cards.append(CardMatch(
            value=" ".join(digits[i:i+4] for i in range(0, len(digits), 4)),
            brand=bin_rec.get("brand") or card_brand_by_prefix(digits),
            card_type=bin_rec.get("type", ""),
            category=bin_rec.get("category", ""),
            issuer=bin_rec.get("issuer", ""),
            issuer_country=bin_rec.get("country", ""),
            bin_code=digits[:8] if len(digits) >= 8 else digits[:6],
            luhn=luhn,
            raw=raw,
            source_file=str(path),
            context=snippet(text, m.start(), m.end()),
        ))

    for m in PHONE_RE.finditer(text):
        raw = m.group(0)
        has_plus = raw.lstrip().startswith("+")
        has_sep = any(c in raw for c in " -()")
        digits = re.sub(r"\D", "", raw)
        if not 10 <= len(digits) <= 15:
            continue
        if not (has_plus or has_sep or detect_country_full(digits)[0]):
            continue
        key = ("p", digits)
        if key in seen:
            continue
        seen.add(key)
        _, country_name = detect_country_full(digits)
        phones.append(PhoneMatch(
            value=format_phone(digits),
            country=country_name or "Неизвестно",
            raw=raw,
            source_file=str(path),
            context=snippet(text, m.start(), m.end()),
        ))

    log.info("%s — карт: %d, телефонов: %d", path, len(cards), len(phones))
    return cards, phones


# ===========================================================================
# 7. ПАРСИНГ TELEGRAM — АККАУНТЫ И СООБЩЕНИЯ
# ===========================================================================
def _flatten_tg_text(raw) -> str:
    if raw is None:
        return ""
    if isinstance(raw, str):
        return raw
    if isinstance(raw, list):
        parts = []
        for item in raw:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict):
                parts.append(item.get("text") or item.get("href") or "")
        return "".join(parts)
    return str(raw)


def parse_telegram_export(path: Path) -> list:
    """Извлекает уникальные учётные записи."""
    try:
        data = json.loads(read_text_auto(path))
    except Exception as e:
        log.error("Не удалось прочитать JSON %s: %s", path, e)
        return []

    accounts: dict[str, TgAccount] = {}

    def add_user(u: dict, chat_name: str):
        if not isinstance(u, dict):
            return
        uid = str(u.get("id") or u.get("user_id") or "")
        if not uid:
            return
        acc = accounts.setdefault(uid, TgAccount(
            user_id=uid,
            username=u.get("username") or "",
            first_name=u.get("first_name") or "",
            last_name=u.get("last_name") or "",
            phone=u.get("phone_number") or u.get("phone") or "",
            is_bot=str(bool(u.get("is_bot"))) if "is_bot" in u else "",
            is_premium=str(bool(u.get("is_premium"))) if "is_premium" in u else "",
            source_chat=chat_name,
            source_file=str(path),
        ))
        for k in ("username", "first_name", "last_name", "phone"):
            if not getattr(acc, k) and u.get(k):
                setattr(acc, k, u.get(k))

    def walk(obj, chat_name=""):
        if isinstance(obj, dict):
            if "name" in obj and ("messages" in obj or "type" in obj):
                chat_name = obj.get("name", chat_name)
            if "personal_information" in obj:
                pi = obj["personal_information"]
                if isinstance(pi, dict):
                    add_user(pi, chat_name or "Личный профиль")
            if obj.get("type") == "message":
                if isinstance(obj.get("from"), dict):
                    add_user(obj["from"], chat_name)
                if isinstance(obj.get("from_id"), str):
                    uid = obj["from_id"].replace("user", "").strip()
                    if uid.isdigit():
                        accounts.setdefault(uid, TgAccount(
                            user_id=uid, username="",
                            first_name=obj.get("from", ""), last_name="",
                            phone="", is_bot="", is_premium="",
                            source_chat=chat_name, source_file=str(path),
                        ))
            for v in obj.values():
                walk(v, chat_name)
        elif isinstance(obj, list):
            for v in obj:
                walk(v, chat_name)

    walk(data)
    log.info("Telegram-экспорт %s — учётных записей: %d", path, len(accounts))
    return list(accounts.values())


def parse_telegram_messages(path: Path) -> list:
    """Извлекает все сообщения с датой, временем и текстом."""
    try:
        data = json.loads(read_text_auto(path))
    except Exception as e:
        log.error("Не удалось прочитать JSON %s: %s", path, e)
        return []

    messages: list[TgMessage] = []

    def walk(node, chat_name="", chat_type=""):
        if isinstance(node, dict):
            if "messages" in node and isinstance(node["messages"], list):
                chat_name = node.get("name", chat_name)
                chat_type = node.get("type", chat_type)
                for msg in node["messages"]:
                    if isinstance(msg, dict) and msg.get("type") in ("message", "service"):
                        rec = _build_message(msg, chat_name, chat_type, path)
                        if rec:
                            messages.append(rec)
            for v in node.values():
                walk(v, chat_name, chat_type)
        elif isinstance(node, list):
            for v in node:
                walk(v, chat_name, chat_type)

    walk(data)
    messages.sort(key=lambda m: (m.datetime_iso, m.msg_id))
    log.info("Telegram-экспорт %s — сообщений: %d", path, len(messages))
    return messages


def _build_message(msg: dict, chat_name: str, chat_type: str, path: Path):
    dt_iso = msg.get("date") or ""
    date_part, time_part = "", ""
    if dt_iso:
        parts = dt_iso.split("T")
        date_part = parts[0]
        if len(parts) > 1:
            time_part = parts[1]

    sender_id = ""
    sender_name = msg.get("from") or ""
    sender_username = ""
    from_id = msg.get("from_id") or ""
    if isinstance(from_id, str):
        sender_id = (from_id
                     .replace("user", "")
                     .replace("channel", "")
                     .replace("chat", "")
                     .strip())
    from_obj = msg.get("from")
    if isinstance(from_obj, dict):
        sender_id = str(from_obj.get("id", sender_id))
        sender_name = from_obj.get("first_name", sender_name)
        sender_username = from_obj.get("username", "")

    reply_to = str(msg["reply_to_message_id"]) if msg.get("reply_to_message_id") else ""
    edited = ""
    if msg.get("edited"):
        edited = str(msg["edited"])
    elif msg.get("edited_unixtime"):
        edited = str(msg["edited_unixtime"])

    raw_type = msg.get("type", "message")
    media_type = msg.get("media_type", "") or ""
    text = _flatten_tg_text(msg.get("text"))

    if not text:
        if media_type:
            text = f"[{media_type}]"
        elif msg.get("sticker"):
            text = "[стикер]"
        elif msg.get("photo"):
            text = "[фото]"
        elif msg.get("file"):
            text = "[файл]"

    if raw_type == "service" and not text:
        action = msg.get("action", "")
        text = f"[{action}]" if action else "[service]"

    msg_type = media_type or ("service" if raw_type == "service" else "text")

    return TgMessage(
        msg_id=str(msg.get("id", "")),
        chat_name=chat_name or "—",
        chat_type=chat_type or "—",
        sender_id=sender_id,
        sender_name=sender_name,
        sender_username=sender_username,
        date=date_part,
        time=time_part,
        datetime_iso=dt_iso or date_part,
        msg_type=msg_type,
        text=text,
        reply_to=reply_to,
        edited=edited,
        source_file=str(path),
    )


def is_telegram_export(path: Path) -> bool:
    if path.suffix.lower() != ".json":
        return False
    try:
        data = json.loads(read_text_auto(path))
    except Exception:
        return False
    if not isinstance(data, dict):
        return False
    return bool({"personal_information", "chats", "about", "contacts"} & set(data.keys()))


# ===========================================================================
# 8. АГРЕГАЦИЯ УНИКАЛЬНЫХ ИДЕНТИФИКАТОРОВ
# ===========================================================================
def aggregate_cards(cards: list) -> list:
    index: dict[str, UniqueCard] = {}
    for c in cards:
        digits = re.sub(r"\D", "", c.value)
        if digits in index:
            u = index[digits]
            u.occurrences += 1
            u.files.add(c.source_file)
        else:
            index[digits] = UniqueCard(
                value=c.value, digits=digits, bin_code=c.bin_code,
                brand=c.brand, card_type=c.card_type, issuer=c.issuer,
                country=c.issuer_country, occurrences=1,
                files={c.source_file},
            )
    result = sorted(index.values(), key=lambda u: (-u.occurrences, u.digits))
    log.info("Уникальных карт: %d", len(result))
    return result


def aggregate_phones(phones: list) -> list:
    index: dict[str, UniquePhone] = {}
    for p in phones:
        digits = re.sub(r"\D", "", p.value)
        iso, name = detect_country_full(digits)
        if digits in index:
            u = index[digits]
            u.occurrences += 1
            u.files.add(p.source_file)
        else:
            index[digits] = UniquePhone(
                value=p.value, digits=digits,
                country_iso=iso or "—",
                country_name=name or p.country,
                occurrences=1, files={p.source_file},
            )
    result = sorted(index.values(), key=lambda u: (-u.occurrences, u.digits))
    log.info("Уникальных телефонов: %d", len(result))
    return result


# ===========================================================================
# 9. ОБХОД ФАЙЛОВ
# ===========================================================================
def collect_files(paths: Iterable[str]) -> list:
    files: list[Path] = []
    for raw in paths:
        p = Path(raw)
        if p.is_dir():
            files.extend(f for f in p.rglob("*")
                         if f.is_file() and f.suffix.lower() in SUPPORTED_EXTS)
        elif p.is_file():
            files.append(p)
        else:
            log.warning("Путь не найден: %s", p)
    return files


# ===========================================================================
# 10. ЗАПИСЬ EXCEL
# ===========================================================================
HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUMMARY_FILL = PatternFill("solid", fgColor="EDEDED")
SUMMARY_LABEL_FONT = Font(bold=True)


def _style_header(ws, headers: list):
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _autosize(ws, widths: list):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_sheet(ws, headers: list, rows: list, widths: list,
                 wrap_col: Optional[int] = None):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_header(ws, headers)
    _autosize(ws, widths)
    if wrap_col:
        for row in ws.iter_rows(min_row=2, min_col=wrap_col, max_col=wrap_col):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_excel(
    cards: list,
    phones: list,
    accounts: list,
    messages: list,
    unique_cards: list,
    unique_phones: list,
    output_path: Path,
    processed_files: list,
) -> None:
    wb = Workbook()

    # ----------------------------------------------------------------------
    # Лист 1: СВОДКА
    # ----------------------------------------------------------------------
    ws_sum = wb.active
    ws_sum.title = "Сводка"

    total_cards_occ = sum(u.occurrences for u in unique_cards)
    total_phones_occ = sum(u.occurrences for u in unique_phones)
    chats = {m.chat_name for m in messages if m.chat_name and m.chat_name != "—"}
    senders = {m.sender_id for m in messages if m.sender_id}
    date_min = min((m.date for m in messages if m.date), default="—")
    date_max = max((m.date for m in messages if m.date), default="—")

    # Страны телефонных номеров
    phone_countries = {}
    for u in unique_phones:
        key = u.country_name or "Неизвестно"
        phone_countries[key] = phone_countries.get(key, 0) + 1

    # Бренды карт
    card_brands = {}
    for u in unique_cards:
        key = u.brand or "Неизвестно"
        card_brands[key] = card_brands.get(key, 0) + 1

    summary = [
        ("Общие показатели", ""),
        ("Обработано файлов", len(processed_files)),
        ("Дата формирования отчёта", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("Карты", ""),
        ("Всего вхождений (с дублями)", total_cards_occ),
        ("Уникальных карт", len(unique_cards)),
        ("Из них прошли Луна", sum(1 for c in cards if c.luhn)),
        ("", ""),
        ("Телефоны", ""),
        ("Всего вхождений (с дублями)", total_phones_occ),
        ("Уникальных номеров", len(unique_phones)),
        ("", ""),
        ("Telegram", ""),
        ("Учётных записей", len(accounts)),
        ("Сообщений (переписок)", len(messages)),
        ("Уникальных чатов", len(chats)),
        ("Уникальных отправителей", len(senders)),
        ("Период переписок", f"{date_min} — {date_max}" if date_min != "—" else "—"),
        ("", ""),
        ("Разбивка карт по брендам", ""),
    ]
    for brand, cnt in sorted(card_brands.items(), key=lambda x: -x[1]):
        summary.append((f"  {brand}", cnt))

    summary.append(("", ""))
    summary.append(("Разбивка телефонов по странам", ""))
    for country, cnt in sorted(phone_countries.items(), key=lambda x: -x[1]):
        summary.append((f"  {country}", cnt))

    for label, value in summary:
        ws_sum.append([label, value])

    # Стилизация сводки
    ws_sum.column_dimensions["A"].width = 42
    ws_sum.column_dimensions["B"].width = 32

    for row_idx, (label, value) in enumerate(summary, start=1):
        cell_a = ws_sum.cell(row=row_idx, column=1)
        cell_b = ws_sum.cell(row=row_idx, column=2)

        # Заголовочные строки
        if label and not value and (
            label.startswith("Общие") or label in {
                "Карты", "Телефоны", "Telegram",
                "Разбивка карт по брендам", "Разбивка телефонов по странам"
            }
        ):
            cell_a.fill = HEADER_FILL
            cell_a.font = HEADER_FONT
            cell_b.fill = HEADER_FILL
        # Метки
        elif label and value != "":
            cell_a.font = SUMMARY_LABEL_FONT

        cell_b.alignment = Alignment(horizontal="left")

    # ----------------------------------------------------------------------
    # Лист 2: Карты (все вхождения)
    # ----------------------------------------------------------------------
    ws_cards = wb.create_sheet("Карты")
    _write_sheet(
        ws_cards,
        ["№", "Номер карты", "BIN", "Бренд", "Тип", "Категория",
         "Банк-эмитент", "Страна эмитента", "Луна",
         "Исходный текст", "Файл", "Контекст"],
        [
            [i, c.value, c.bin_code, c.brand, c.card_type, c.category,
             c.issuer, c.issuer_country, "OK" if c.luhn else "—",
             c.raw, c.source_file, c.context]
            for i, c in enumerate(cards, 1)
        ],
        [6, 26, 12, 16, 12, 14, 32, 20, 8, 24, 45, 70],
    )

    # ----------------------------------------------------------------------
    # Лист 3: Телефоны (все вхождения)
    # ----------------------------------------------------------------------
    ws_phones = wb.create_sheet("Телефоны")
    _write_sheet(
        ws_phones,
        ["№", "Номер", "Страна", "Исходный текст", "Файл", "Контекст"],
        [
            [i, p.value, p.country, p.raw, p.source_file, p.context]
            for i, p in enumerate(phones, 1)
        ],
        [6, 26, 20, 24, 45, 70],
    )

    # ----------------------------------------------------------------------
    # Лист 4: ID Карты (уникальные)
    # ----------------------------------------------------------------------
    ws_uc = wb.create_sheet("ID Карты")
    _write_sheet(
        ws_uc,
        ["№", "Номер карты", "Цифры (без пробелов)", "BIN",
         "Бренд", "Тип", "Банк-эмитент", "Страна",
         "Встречается раз", "Файлы"],
        [
            [i, u.value, u.digits, u.bin_code, u.brand, u.card_type,
             u.issuer, u.country, u.occurrences, "; ".join(sorted(u.files))]
            for i, u in enumerate(unique_cards, 1)
        ],
        [6, 26, 22, 12, 16, 12, 32, 20, 14, 60],
    )

    # ----------------------------------------------------------------------
    # Лист 5: ID Телефоны (уникальные)
    # ----------------------------------------------------------------------
    ws_up = wb.create_sheet("ID Телефоны")
    _write_sheet(
        ws_up,
        ["№", "Номер", "Цифры (без пробелов)",
         "Код страны (ISO)", "Страна",
         "Встречается раз", "Файлы"],
        [
            [i, u.value, u.digits, u.country_iso, u.country_name,
             u.occurrences, "; ".join(sorted(u.files))]
            for i, u in enumerate(unique_phones, 1)
        ],
        [6, 26, 22, 16, 22, 14, 60],
    )

    # ----------------------------------------------------------------------
    # Лист 6: TG Учётные записи
    # ----------------------------------------------------------------------
    ws_acc = wb.create_sheet("TG Учётные записи")
    _write_sheet(
        ws_acc,
        ["№", "User ID", "Username", "Имя", "Фамилия", "Телефон",
         "Бот", "Premium", "Чат", "Файл"],
        [
            [i, a.user_id,
             f"@{a.username}" if a.username else "",
             a.first_name, a.last_name, a.phone,
             a.is_bot, a.is_premium, a.source_chat, a.source_file]
            for i, a in enumerate(accounts, 1)
        ],
        [6, 16, 22, 20, 20, 20, 8, 10, 30, 45],
    )

    # ----------------------------------------------------------------------
    # Лист 7: Переписки
    # ----------------------------------------------------------------------
    ws_msg = wb.create_sheet("Переписки")
    _write_sheet(
        ws_msg,
        ["№", "Дата", "Время", "Дата/время (полное)",
         "Чат", "Тип чата",
         "ID отправителя", "Имя отправителя", "Username",
         "Тип сообщения", "Текст сообщения",
         "Ответ на #", "Редактировано", "Файл"],
        [
            [i, m.date, m.time, m.datetime_iso,
             m.chat_name, m.chat_type,
             m.sender_id, m.sender_name,
             f"@{m.sender_username}" if m.sender_username else "",
             m.msg_type, m.text,
             m.reply_to, m.edited, m.source_file]
            for i, m in enumerate(messages, 1)
        ],
        [6, 12, 10, 20, 28, 16, 16, 22, 22, 14, 80, 12, 20, 40],
        wrap_col=11,
    )

    wb.save(output_path)
    log.info("Результат сохранён: %s", output_path)


# ===========================================================================
# 11. CLI
# ===========================================================================
def main() -> int:
    ap = argparse.ArgumentParser(
        description="Извлекает карты, телефоны, TG-аккаунты и переписки в Excel."
    )
    ap.add_argument("paths", nargs="+", help="Файлы или папки для обработки")
    ap.add_argument("-o", "--output", default="result.xlsx",
                    help="Итоговый Excel (по умолчанию result.xlsx)")
    ap.add_argument("--bin-db", default="bin-list-data.csv",
                    help="Путь к CSV-справочнику BIN")
    ap.add_argument("--chat", default=None,
                    help="Оставить только сообщения из чата, содержащего эту подстроку")
    ap.add_argument("--from", dest="date_from", default=None,
                    help="Только сообщения с даты YYYY-MM-DD")
    ap.add_argument("--to", dest="date_to", default=None,
                    help="Только сообщения по дату YYYY-MM-DD")
    args = ap.parse_args()

    bindb = BinDatabase(args.bin_db)
    if not bindb.loaded:
        log.warning(
            "BIN-справочник не загружен. Бренд будет по префиксу, "
            "эмитент — недоступен. Скачайте:\n"
            "  curl -L -o bin-list-data.csv "
            "https://raw.githubusercontent.com/venelinkochev/bin-list-data/"
            "master/bin-list-data.csv"
        )

    files = collect_files(args.paths)
    if not files:
        log.error("Не найдено ни одного файла.")
        return 1
    log.info("Файлов к обработке: %d", len(files))

    all_cards, all_phones = [], []
    all_accounts, all_messages = [], []

    for f in files:
        if is_telegram_export(f):
            log.info("Обнаружен экспорт Telegram: %s", f)
            all_accounts.extend(parse_telegram_export(f))
            all_messages.extend(parse_telegram_messages(f))
            cs, ps = process_file(f, bindb)
            all_cards.extend(cs)
            all_phones.extend(ps)
        else:
            cs, ps = process_file(f, bindb)
            all_cards.extend(cs)
            all_phones.extend(ps)

    # Фильтры по сообщениям
    if args.chat:
        needle = args.chat.lower()
        all_messages = [m for m in all_messages if needle in m.chat_name.lower()]
    if args.date_from:
        all_messages = [m for m in all_messages if m.date >= args.date_from]
    if args.date_to:
        all_messages = [m for m in all_messages if m.date <= args.date_to]

    unique_cards = aggregate_cards(all_cards)
    unique_phones = aggregate_phones(all_phones)

    write_excel(
        all_cards, all_phones, all_accounts, all_messages,
        unique_cards, unique_phones,
        Path(args.output), files,
    )

    log.info(
        "Итого: карт %d (уникальных %d), телефонов %d (уникальных %d), "
        "TG-аккаунтов %d, сообщений %d",
        len(all_cards), len(unique_cards),
        len(all_phones), len(unique_phones),
        len(all_accounts), len(all_messages),
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())